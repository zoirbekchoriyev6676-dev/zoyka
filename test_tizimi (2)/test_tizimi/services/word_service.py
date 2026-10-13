import re
import io
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

class WordService:
    @staticmethod
    def parse_docx(file_bytes: bytes) -> dict:
        """
        Word (.docx) faylini xavfsiz o'qib, savollar, variantlar va to'g'ri javoblarni ajratish.
        Server hech qachon yiqilmaydi — xatolar har bir savol darajasida qayd etiladi.
        """
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
        except Exception as e:
            return {
                "success": False,
                "error": f"Word faylini ochishda xatolik yuz berdi: {str(e)}",
                "questions": [],
                "total_found": 0,
                "valid_count": 0,
                "error_count": 0,
                "test_title": "Noma'lum Test"
            }

        # Matn qatorlarini yig'ish (paragraflar va jadvallardan)
        raw_lines = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                for line in text.split("\n"):
                    l_clean = line.strip()
                    if l_clean:
                        raw_lines.append(l_clean)

        # Jadvallardan ham qatorlarni qo'shish
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    ctext = cell.text.strip()
                    if ctext:
                        for line in ctext.split("\n"):
                            l_clean = line.strip()
                            if l_clean:
                                raw_lines.append(l_clean)

        if not raw_lines:
            return {
                "success": False,
                "error": "Word fayli bo'sh yoki undan matn o'qib bo'lmadi",
                "questions": [],
                "total_found": 0,
                "valid_count": 0,
                "error_count": 0,
                "test_title": "Bo'sh Test"
            }

        test_title = "Word orqali yuklangan test"
        first_line = raw_lines[0]
        if re.match(r"^(test\s+nomi|test)\s*:\s*", first_line, re.IGNORECASE):
            test_title = re.sub(r"^(test\s+nomi|test)\s*:\s*", "", first_line, flags=re.IGNORECASE).strip()
            raw_lines = raw_lines[1:]

        # Savollar bloklarini ajratish
        parsed_questions = []
        current_q = None

        opt_pattern = re.compile(r"^([A-Da-d])[\)\.\:]\s*(.+)$")
        ans_pattern = re.compile(r"^(javob|to['`’]g['`’]ri\s*javob|answer|variant)\s*[:=\-]?\s*([A-Da-d])", re.IGNORECASE)
        topic_pattern = re.compile(r"^(mavzu|topic)\s*[:=\-]?\s*(.+)$", re.IGNORECASE)
        diff_pattern = re.compile(r"^(qiyinlik|difficulty)\s*[:=\-]?\s*(.+)$", re.IGNORECASE)
        points_pattern = re.compile(r"^(ball|points?)\s*[:=\-]?\s*(\d+)$", re.IGNORECASE)
        exp_pattern = re.compile(r"^(izoh|tushuntirish|explanation)\s*[:=\-]?\s*(.+)$", re.IGNORECASE)
        new_q_pattern = re.compile(r"^(\d+[\.\)]\s*|savol\s*\d*[\:\.]\s*)", re.IGNORECASE)

        def finalize_question(q):
            if not q:
                return
            # Tekshirish va xatolik belgilash
            has_error = False
            error_msg = None

            if len(q["options"]) < 2:
                has_error = True
                error_msg = f"{q['temp_id']}-savolda variantlar yetarli emas (kamida 2 ta bo'lishi kerak)"
            else:
                correct_count = sum(1 for opt in q["options"] if opt.get("is_correct") == 1)
                if correct_count == 0:
                    has_error = True
                    error_msg = f"{q['temp_id']}-savolda to'g'ri javob belgilanmagan"
                elif correct_count > 1:
                    has_error = True
                    error_msg = f"{q['temp_id']}-savolda bir nechta to'g'ri javob ko'rsatilgan"

            q["has_error"] = has_error
            q["error_message"] = error_msg
            parsed_questions.append(q)

        temp_id_counter = 1

        for line in raw_lines:
            line_str = line.strip()
            if not line_str:
                continue

            # Yangi savol boshlanishi (masalan, "1. 2+2=?" yoki "Savol 1:")
            # Yoki variantlar to'lib bo'lgandan keyin yangi oddiy qator kelsa
            is_new_q_number = bool(new_q_pattern.match(line_str))
            
            # Agar mavjud savolda variantlar bor bo'lsa va yangi qator variant/meta bo'lmasa, demak yangi savol
            is_new_q_start = False
            if current_q and len(current_q["options"]) >= 2:
                if (not opt_pattern.match(line_str) and 
                    not ans_pattern.match(line_str) and 
                    not topic_pattern.match(line_str) and 
                    not diff_pattern.match(line_str) and 
                    not points_pattern.match(line_str) and 
                    not exp_pattern.match(line_str)):
                    is_new_q_start = True

            if is_new_q_number or is_new_q_start or current_q is None:
                if current_q is not None:
                    finalize_question(current_q)
                
                clean_q_text = new_q_pattern.sub("", line_str).strip()
                if not clean_q_text:
                    clean_q_text = line_str
                
                current_q = {
                    "temp_id": temp_id_counter,
                    "question_text": clean_q_text,
                    "topic_name": "Umumiy",
                    "difficulty": "medium",
                    "points": 1,
                    "explanation": "",
                    "options": [],
                    "correct_letter": None
                }
                temp_id_counter += 1
                continue

            # Variantmi?
            opt_match = opt_pattern.match(line_str)
            if opt_match and current_q is not None:
                letter = opt_match.group(1).upper()
                opt_text = opt_match.group(2).strip()
                is_corr = 1 if (current_q["correct_letter"] and current_q["correct_letter"] == letter) else 0
                current_q["options"].append({
                    "letter": letter,
                    "option_text": opt_text,
                    "is_correct": is_corr
                })
                continue

            # Javobmi?
            ans_match = ans_pattern.match(line_str)
            if ans_match and current_q is not None:
                correct_letter = ans_match.group(2).upper()
                current_q["correct_letter"] = correct_letter
                # Oldingi kiritilgan variantlar ichidan to'g'risini belgilash
                for opt in current_q["options"]:
                    if opt.get("letter") == correct_letter:
                        opt["is_correct"] = 1
                    else:
                        opt["is_correct"] = 0
                continue

            # Mavzumi?
            topic_match = topic_pattern.match(line_str)
            if topic_match and current_q is not None:
                current_q["topic_name"] = topic_match.group(2).strip()
                continue

            # Qiyinlikmi?
            diff_match = diff_pattern.match(line_str)
            if diff_match and current_q is not None:
                d_val = diff_match.group(2).strip().lower()
                if "oson" in d_val or "easy" in d_val:
                    current_q["difficulty"] = "easy"
                elif "qiyin" in d_val or "hard" in d_val:
                    current_q["difficulty"] = "hard"
                else:
                    current_q["difficulty"] = "medium"
                continue

            # Ballmi?
            pts_match = points_pattern.match(line_str)
            if pts_match and current_q is not None:
                try:
                    current_q["points"] = max(1, int(pts_match.group(2)))
                except ValueError:
                    pass
                continue

            # Izohmi?
            exp_match = exp_pattern.match(line_str)
            if exp_match and current_q is not None:
                current_q["explanation"] = exp_match.group(2).strip()
                continue

            # Agar yuqoridagilar bo'lmasa, savol matniga qo'shish
            if current_q is not None and not current_q["options"]:
                current_q["question_text"] += " " + line_str

        if current_q is not None:
            finalize_question(current_q)

        valid_count = sum(1 for q in parsed_questions if not q["has_error"])
        error_count = sum(1 for q in parsed_questions if q["has_error"])

        return {
            "success": True,
            "test_title": test_title,
            "total_found": len(parsed_questions),
            "valid_count": valid_count,
            "error_count": error_count,
            "questions": parsed_questions
        }

    @staticmethod
    def generate_docx_export(test_data: dict, mode: str = "teacher") -> io.BytesIO:
        """
        Mavjud testni Word (.docx) formatida eksport qilish.
        mode: 'student' (javoblarsiz) yoki 'teacher' (to'g'ri javoblar va izohlar bilan)
        """
        doc = docx.Document()

        # Sahifa parametrlari
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Sarlavha
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.add_run(test_data.get("title", "Imtihon Testi"))
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_run.font.name = "Arial"
        title_run.font.color.rgb = RGBColor(30, 27, 75)

        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        mode_label = "O'QITUVCHI NUSXASI (JAVOBLAR BILAN)" if mode == "teacher" else "O'QUVCHI NUSXASI"
        sub_run = subtitle.add_run(mode_label)
        sub_run.font.size = Pt(10)
        sub_run.font.bold = True
        sub_run.font.color.rgb = RGBColor(99, 102, 241) if mode == "teacher" else RGBColor(100, 116, 139)

        # Meta Ma'lumotlar Jadvali
        table = doc.add_table(rows=3, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        meta = [
            ("Fan:", test_data.get("subject_name", "-"), "Guruh:", test_data.get("group_name", "-")),
            ("O'qituvchi:", test_data.get("teacher_name", "-"), "Davomiyligi:", f"{test_data.get('duration_minutes', 30)} daqiqa"),
            ("O'tish bali:", f"{test_data.get('passing_score', 60)}%", "Savollar soni:", f"{len(test_data.get('questions', []))} ta")
        ]

        for i, (k1, v1, k2, v2) in enumerate(meta):
            row_cells = table.rows[i].cells
            row_cells[0].text = f"{k1} {v1}"
            row_cells[1].text = f"{k2} {v2}"
            for c in row_cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        r.font.name = "Arial"

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # Savollar ro'yxati
        letters = ["A", "B", "C", "D", "E", "F"]
        for idx, q in enumerate(test_data.get("questions", []), 1):
            qp = doc.add_paragraph()
            qp.paragraph_format.space_before = Pt(8)
            qp.paragraph_format.space_after = Pt(3)

            q_num_run = qp.add_run(f"{idx}. ")
            q_num_run.font.bold = True
            q_num_run.font.size = Pt(11)
            q_num_run.font.name = "Arial"

            q_text_run = qp.add_run(q.get("question_text", ""))
            q_text_run.font.bold = True
            q_text_run.font.size = Pt(11)
            q_text_run.font.name = "Arial"

            # Metadata belgisi (faqat o'qituvchi rejimida yoki kichik izoh sifatida)
            if mode == "teacher":
                diff_uz = {"easy": "Oson", "medium": "O'rta", "hard": "Qiyin"}.get(q.get("difficulty"), "O'rta")
                meta_run = qp.add_run(f"  [{diff_uz} • {q.get('points', 1)} ball • {q.get('topic_name', '')}]")
                meta_run.font.size = Pt(8.5)
                meta_run.font.italic = True
                meta_run.font.color.rgb = RGBColor(148, 163, 184)

            # Variantlar
            correct_letter = None
            for o_idx, opt in enumerate(q.get("options", [])):
                letter = letters[o_idx] if o_idx < len(letters) else f"({o_idx+1})"
                is_correct = bool(opt.get("is_correct"))
                if is_correct:
                    correct_letter = letter

                op = doc.add_paragraph()
                op.paragraph_format.left_indent = Inches(0.25)
                op.paragraph_format.space_after = Pt(2)
                op.paragraph_format.space_before = Pt(0)

                o_run = op.add_run(f"{letter}) {opt.get('option_text', '')}")
                o_run.font.size = Pt(10)
                o_run.font.name = "Arial"

                if mode == "teacher" and is_correct:
                    o_run.font.bold = True
                    o_run.font.color.rgb = RGBColor(22, 101, 52)  # To'q yashil
                    chk_run = op.add_run("  ✓ [TO'G'RI JAVOB]")
                    chk_run.font.bold = True
                    chk_run.font.size = Pt(8.5)
                    chk_run.font.color.rgb = RGBColor(22, 101, 52)

            # O'qituvchi rejimida to'g'ri javob va izoh xulosasi
            if mode == "teacher":
                ans_p = doc.add_paragraph()
                ans_p.paragraph_format.left_indent = Inches(0.25)
                ans_p.paragraph_format.space_after = Pt(6)
                
                if correct_letter:
                    ans_run = ans_p.add_run(f"To'g'ri javob: {correct_letter}")
                    ans_run.font.bold = True
                    ans_run.font.size = Pt(9.5)
                    ans_run.font.color.rgb = RGBColor(79, 70, 229)

                explanation = q.get("explanation")
                if explanation:
                    exp_run = ans_p.add_run(f"  |  Izoh: {explanation}")
                    exp_run.font.italic = True
                    exp_run.font.size = Pt(9)
                    exp_run.font.color.rgb = RGBColor(100, 116, 139)

        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        return stream
